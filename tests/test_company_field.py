"""
Tests for Candidate Company / Organization Name feature.

Covers:
1. Entry form rendering optional company field
2. Storing and normalizing company_name on candidate entry (whitespace trimming, max length)
3. Omitting or leaving company empty results in None
4. Creating StudentAttempt saves company_name
5. Resuming existing in-progress attempt preserves original company_name
6. Admin attempts page filtering by company
7. Admin attempts dropdown listing distinct companies
8. Admin attempts table displaying company badge/name when present and clean layout when None
9. Export results Excel includes 'Company / Organization' column with correct values or '-'
10. Export selected results Excel includes 'Company / Organization' column
"""
import io
import uuid
import openpyxl
import pytest
from app.extensions import db
from app.models import Exam, Question, QuestionOption, StudentAttempt


def _create_test_exam(title="Company Test Exam", code="COMPTEST01"):
    exam = Exam(
        title=title,
        exam_code=code,
        duration_minutes=45,
        passing_type="percentage",
        passing_value=50.0,
    )
    db.session.add(exam)
    db.session.flush()

    q = Question(
        exam_id=exam.id,
        question_text="Sample Q",
        question_type="mcq",
        marks=2.0,
    )
    db.session.add(q)
    db.session.flush()

    opt1 = QuestionOption(question_id=q.id, option_text="Option A", option_order=0)
    opt2 = QuestionOption(question_id=q.id, option_text="Option B", option_order=1)
    db.session.add_all([opt1, opt2])
    db.session.flush()
    q.correct_option_id = opt1.id
    db.session.commit()
    return exam


def test_entry_form_renders_company_field(client, app):
    """GET /exam/<code > renders the optional company_name input."""
    with app.app_context():
        exam = _create_test_exam("Entry Form Exam", "COMPENTRY1")
        code = exam.exam_code

    resp = client.get(f"/exam/{code}")
    assert resp.status_code == 200
    html = resp.data.decode("utf-8")
    assert 'name="company_name"' in html
    assert "Company / Organization Name" in html
    assert "(Optional)" in html


def test_entry_post_normalizes_company(client, app):
    """POST /exam/<code> normalizes company name (whitespace trimmed, collapsed, max 150)."""
    with app.app_context():
        exam = _create_test_exam("Normalize Exam", "COMPNORM01")
        code = exam.exam_code

    # 1. Normal company with extra spaces
    resp = client.post(
        f"/exam/{code}",
        data={
            "student_name": "John Doe",
            "student_email": "john@example.com",
            "company_name": "   Acme   Widgets   Pvt   Ltd   ",
        },
        follow_redirects=False,
    )
    assert resp.status_code == 302
    with client.session_transaction() as sess:
        assert sess.get("company_name") == "Acme Widgets Pvt Ltd"

    # 2. Blank / whitespace company stored as None
    client.post(
        f"/exam/{code}",
        data={
            "student_name": "Jane Doe",
            "student_email": "jane@example.com",
            "company_name": "   ",
        },
        follow_redirects=False,
    )
    with client.session_transaction() as sess:
        assert sess.get("company_name") is None

    # 3. Truncated to 150 characters
    long_company = "A" * 200
    client.post(
        f"/exam/{code}",
        data={
            "student_name": "Long Name",
            "student_email": "long@example.com",
            "company_name": long_company,
        },
        follow_redirects=False,
    )
    with client.session_transaction() as sess:
        assert len(sess.get("company_name")) == 150


def test_start_exam_creates_attempt_with_company(client, app):
    """Starting exam creates a StudentAttempt with the candidate's company."""
    with app.app_context():
        exam = _create_test_exam("Start Exam Test", "COMPSTART1")
        code = exam.exam_code

    with client.session_transaction() as sess:
        sess["student_name"] = "Alice Candidate"
        sess["student_email"] = "alice@acme.com"
        sess["company_name"] = "Acme Corp"

    resp = client.post(f"/exam/{code}/start", follow_redirects=False)
    assert resp.status_code == 302

    with app.app_context():
        attempt = StudentAttempt.query.filter_by(student_email="alice@acme.com").first()
        assert attempt is not None
        assert attempt.company_name == "Acme Corp"
        assert attempt.student_name == "Alice Candidate"

    # Verify company_name is popped from session
    with client.session_transaction() as sess:
        assert "company_name" not in sess


def test_start_exam_resume_preserves_company(client, app):
    """Resuming an in-progress attempt preserves original company_name."""
    with app.app_context():
        exam = _create_test_exam("Resume Exam", "COMPRESUME1")
        token = str(uuid.uuid4())
        existing = StudentAttempt(
            exam_id=exam.id,
            student_name="Bob Existing",
            student_email="bob@original.com",
            company_name="Original Co",
            attempt_token=token,
            status="in_progress",
        )
        db.session.add(existing)
        db.session.commit()
        code = exam.exam_code

    with client.session_transaction() as sess:
        sess["student_name"] = "Bob Existing"
        sess["student_email"] = "bob@original.com"
        sess["company_name"] = "Different Co"

    resp = client.post(f"/exam/{code}/start", follow_redirects=False)
    assert resp.status_code == 302
    assert f"/attempt/{token}/question/1" in resp.headers["Location"]

    with app.app_context():
        att = StudentAttempt.query.filter_by(attempt_token=token).first()
        # Preserves original company name
        assert att.company_name == "Original Co"


def test_admin_attempts_page_and_filtering(logged_in_client, app):
    """Admin attempts page shows companies dropdown and filters attempts by company."""
    with app.app_context():
        exam = _create_test_exam("Admin Filter Exam", "COMPFILT01")
        exam_id = exam.id

        # Attempt 1: Alpha Tech
        att1 = StudentAttempt(
            exam_id=exam_id,
            student_name="Alpha User",
            student_email="alpha@test.com",
            company_name="Alpha Tech",
            attempt_token=str(uuid.uuid4()),
            status="submitted",
        )
        # Attempt 2: Beta Corp
        att2 = StudentAttempt(
            exam_id=exam_id,
            student_name="Beta User",
            student_email="beta@test.com",
            company_name="Beta Corp",
            attempt_token=str(uuid.uuid4()),
            status="submitted",
        )
        # Attempt 3: No company (NULL)
        att3 = StudentAttempt(
            exam_id=exam_id,
            student_name="Solo User",
            student_email="solo@test.com",
            company_name=None,
            attempt_token=str(uuid.uuid4()),
            status="submitted",
        )
        db.session.add_all([att1, att2, att3])
        db.session.commit()

    # 1. GET without filter: shows all 3, company dropdown has Alpha Tech and Beta Corp
    resp = logged_in_client.get(f"/admin/exams/{exam_id}/attempts")
    assert resp.status_code == 200
    html = resp.data.decode("utf-8")
    assert "Alpha Tech" in html
    assert "Beta Corp" in html
    assert "Solo User" in html
    assert '<option value="Alpha Tech"' in html
    assert '<option value="Beta Corp"' in html

    # 2. GET with company=Alpha Tech: only Alpha User appears
    resp_filtered = logged_in_client.get(f"/admin/exams/{exam_id}/attempts?company=Alpha+Tech")
    assert resp_filtered.status_code == 200
    html_f = resp_filtered.data.decode("utf-8")
    assert "Alpha User" in html_f
    assert "Beta User" not in html_f
    assert "Solo User" not in html_f
    # Clear button should be visible
    assert "Clear" in html_f


def test_export_results_includes_company(logged_in_client, app):
    """Excel export includes 'Company / Organization' column and values."""
    with app.app_context():
        exam = _create_test_exam("Export Exam", "COMPEXPO01")
        exam_id = exam.id

        att1 = StudentAttempt(
            exam_id=exam_id,
            student_name="Export User 1",
            student_email="exp1@test.com",
            company_name="Delta Systems",
            attempt_token=str(uuid.uuid4()),
            status="submitted",
            total_marks_obtained=2.0,
            percentage_score=100.0,
        )
        att2 = StudentAttempt(
            exam_id=exam_id,
            student_name="Export User 2",
            student_email="exp2@test.com",
            company_name=None,
            attempt_token=str(uuid.uuid4()),
            status="submitted",
            total_marks_obtained=0.0,
            percentage_score=0.0,
        )
        db.session.add_all([att1, att2])
        db.session.commit()

    resp = logged_in_client.get(f"/admin/exams/{exam_id}/export-results")
    assert resp.status_code == 200
    assert resp.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    assert "Company / Organization" in headers
    company_col_idx = headers.index("Company / Organization") + 1

    # Row 2 (Delta Systems)
    row2_vals = [cell.value for cell in ws[2]]
    assert row2_vals[company_col_idx - 1] == "Delta Systems"

    # Row 3 (NULL company -> '-')
    row3_vals = [cell.value for cell in ws[3]]
    assert row3_vals[company_col_idx - 1] == "-"


def test_export_filtered_by_company(logged_in_client, app):
    """Excel export respects company filter query parameter."""
    with app.app_context():
        exam = _create_test_exam("Export Filter Exam", "COMPEXPO02")
        exam_id = exam.id

        att1 = StudentAttempt(
            exam_id=exam_id,
            student_name="User A",
            student_email="a@test.com",
            company_name="Gamma Industries",
            attempt_token=str(uuid.uuid4()),
            status="submitted",
        )
        att2 = StudentAttempt(
            exam_id=exam_id,
            student_name="User B",
            student_email="b@test.com",
            company_name="Zeta Labs",
            attempt_token=str(uuid.uuid4()),
            status="submitted",
        )
        db.session.add_all([att1, att2])
        db.session.commit()

    resp = logged_in_client.get(f"/admin/exams/{exam_id}/export-results?company=Gamma+Industries")
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    # Header + 1 data row
    assert len(rows) == 2
    assert rows[1][0] == "User A"
    assert rows[1][2] == "Gamma Industries"


def test_export_selected_results_includes_company(logged_in_client, app):
    """Excel export for selected attempts includes 'Company / Organization' column."""
    with app.app_context():
        exam = _create_test_exam("Export Sel Exam", "COMPSEL01")
        exam_id = exam.id

        att1 = StudentAttempt(
            exam_id=exam_id,
            student_name="Selected User",
            student_email="sel@test.com",
            company_name="Omega Corp",
            attempt_token=str(uuid.uuid4()),
            status="submitted",
        )
        db.session.add(att1)
        db.session.commit()
        att1_id = att1.id

    resp = logged_in_client.post(
        f"/admin/exams/{exam_id}/export-selected-results",
        data={"attempt_ids": [att1_id]},
    )
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][2] == "Company / Organization"
    assert rows[1][0] == "Selected User"
    assert rows[1][2] == "Omega Corp"
