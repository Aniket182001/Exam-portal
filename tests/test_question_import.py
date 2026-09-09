import io
import os
import docx
import pytest

from app.extensions import db
from app.models import Exam, Question, QuestionOption
from app.services.import_parsers import ExcelParser, DocxParser, PdfParser, get_parser


import uuid

def make_test_exam():
    """Create a minimal exam for testing import routes."""
    code = f"IMP_{uuid.uuid4().hex[:8]}"
    exam = Exam(
        title="Import Test Exam",
        exam_code=code,
        duration_minutes=45,
        passing_type="percentage",
        passing_value=50.0,
    )
    db.session.add(exam)
    db.session.commit()
    return Exam.query.filter_by(exam_code=code).first()


def test_excel_parser_with_static_template():
    """
    Verify that the authoritative static Excel upload template
    is parsed accurately by ExcelParser without errors.
    """
    template_path = os.path.join("app", "static", "templates", "exam_upload_template.xlsx")
    assert os.path.exists(template_path), "Static Excel template must exist"

    parser = ExcelParser()
    questions = parser.parse(template_path)

    assert len(questions) == 1
    q = questions[0]
    assert q["question"] == "Which are programming languages?"
    assert len(q["options"]) == 6
    assert q["options"] == ["Python", "Java", "C++", "Ruby", "Go", "Rust"]
    assert q["correct_option_index"] == 0
    assert q["marks"] == 1.0


def test_docx_parser(tmp_path):
    """
    Verify that DocxParser correctly parses structured Word documents.
    """
    doc = docx.Document()
    doc.add_paragraph("Question 1: What is Six Sigma?")
    doc.add_paragraph("A. A martial art technique")
    doc.add_paragraph("B. A data-driven quality methodology")
    doc.add_paragraph("C. A cloud storage solution")
    doc.add_paragraph("Answer: B")
    doc.add_paragraph("Marks: 2.0")

    docx_file = str(tmp_path / "test_exam.docx")
    doc.save(docx_file)

    parser = DocxParser()
    questions = parser.parse(docx_file)

    assert len(questions) == 1
    q = questions[0]
    assert q["question"] == "What is Six Sigma?"
    assert len(q["options"]) == 3
    assert q["options"] == [
        "A martial art technique",
        "A data-driven quality methodology",
        "A cloud storage solution",
    ]
    assert q["correct_option_index"] == 1
    assert q["marks"] == 2.0


def test_get_parser_factory():
    """Verify get_parser returns the appropriate parser class or raises ValueError."""
    assert isinstance(get_parser("sample.xlsx"), ExcelParser)
    assert isinstance(get_parser("sample.docx"), DocxParser)
    assert isinstance(get_parser("sample.pdf"), PdfParser)

    with pytest.raises(ValueError, match="Unsupported file format"):
        get_parser("sample.txt")


def test_upload_and_preview_flow_excel(logged_in_client, app):
    """
    Test uploading the Excel template through the admin web interface,
    verifying clean server-side HTML preview rendering without AI workspace artifacts.
    """
    with app.app_context():
        exam = make_test_exam()
        exam_id = exam.id

    template_path = os.path.join("app", "static", "templates", "exam_upload_template.xlsx")
    with open(template_path, "rb") as f:
        excel_bytes = f.read()

    response = logged_in_client.post(
        f"/admin/exams/{exam_id}/questions/import",
        data={"file": (io.BytesIO(excel_bytes), "upload.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    html = response.get_data(as_text=True)

    # Server-side preview table assertions
    assert "Data Preview" in html
    assert "Parsed Questions (1)" in html
    assert "Which are programming languages?" in html
    assert "Python" in html
    assert "Rust" in html
    assert "Confirm & Save" in html

    # Verify that experimental AI workspace artifacts are absent
    assert "AI Review Workspace" not in html
    assert "Toggle Legacy View" not in html
    assert "Start Review Queue" not in html
    assert "dash-total" not in html


def test_upload_and_preview_flow_docx(logged_in_client, app, tmp_path):
    """
    Test uploading a Word document through the admin web interface
    and rendering in server-side preview.
    """
    with app.app_context():
        exam = make_test_exam()
        exam_id = exam.id

    doc = docx.Document()
    doc.add_paragraph("Question 1: What does DMAIC stand for?")
    doc.add_paragraph("A. Define, Measure, Analyze, Improve, Control")
    doc.add_paragraph("B. Data, Metric, Analysis, Input, Check")
    doc.add_paragraph("Answer: A")
    doc.add_paragraph("Marks: 1.0")

    docx_file = str(tmp_path / "dmaic.docx")
    doc.save(docx_file)

    with open(docx_file, "rb") as f:
        docx_bytes = f.read()

    response = logged_in_client.post(
        f"/admin/exams/{exam_id}/questions/import",
        data={"file": (io.BytesIO(docx_bytes), "dmaic.docx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    html = response.get_data(as_text=True)

    assert "Data Preview" in html
    assert "Parsed Questions (1)" in html
    assert "What does DMAIC stand for?" in html
    assert "Define, Measure, Analyze, Improve, Control" in html


def test_confirm_import_saves_to_database(logged_in_client, app):
    """
    Test that confirming the import persists questions and options
    directly to the database with correct relationships.
    """
    with app.app_context():
        exam = make_test_exam()
        exam_id = exam.id

    template_path = os.path.join("app", "static", "templates", "exam_upload_template.xlsx")
    with open(template_path, "rb") as f:
        excel_bytes = f.read()

    # Step 1: Upload
    upload_res = logged_in_client.post(
        f"/admin/exams/{exam_id}/questions/import",
        data={"file": (io.BytesIO(excel_bytes), "upload.xlsx")},
        content_type="multipart/form-data",
    )
    assert upload_res.status_code == 302
    redirect_url = upload_res.headers["Location"]
    import_id = redirect_url.rstrip("/").split("/")[-1]

    # Step 2: Confirm
    confirm_res = logged_in_client.post(
        f"/admin/exams/{exam_id}/questions/import/confirm",
        data={"import_id": import_id},
        follow_redirects=True,
    )
    assert confirm_res.status_code == 200
    html = confirm_res.get_data(as_text=True)
    assert "Successfully imported 1 questions." in html

    # Step 3: Verify in DB
    with app.app_context():
        created_q = Question.query.filter_by(exam_id=exam_id).first()
        assert created_q is not None
        assert created_q.question_text == "Which are programming languages?"
        assert created_q.marks == 1.0
        assert len(created_q.options) == 6

        # Correct option check (Option A: "Python")
        correct_opt = db.session.get(QuestionOption, created_q.correct_option_id)
        assert correct_opt is not None
        assert correct_opt.option_text == "Python"
        assert correct_opt.option_order == 1


def test_upload_unsupported_file_fails_cleanly(logged_in_client, app):
    """
    Test uploading an unsupported file format fails cleanly with an error flash.
    """
    with app.app_context():
        exam = make_test_exam()
        exam_id = exam.id

    response = logged_in_client.post(
        f"/admin/exams/{exam_id}/questions/import",
        data={"file": (io.BytesIO(b"Not an exam"), "invalid.txt")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert "Error parsing file: Unsupported file format: .txt" in html
